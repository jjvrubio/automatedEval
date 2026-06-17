from docx import Document


def create_cultural_analysis_report():
    doc = Document()

    # Add title
    doc.add_heading("Cultural Analysis Report", level=1)

    # Overview of Cultural Dimensions
    doc.add_heading("Overview of Cultural Dimensions", level=2)
    doc.add_paragraph("- Power Distance: Moderate (Score: 50)")
    doc.add_paragraph("- Individualism: Low (Score: 30)")
    doc.add_paragraph("- Uncertainty Avoidance: High (Score: 80)")

    # Application to Target Region
    doc.add_heading("Application to Target Region", level=2)
    doc.add_paragraph("The region values group harmony and collective decision-making.")

    # Strategic Recommendations
    doc.add_heading("Strategic Recommendations", level=2)
    recommendations = [
        "Emphasize community benefits in marketing messages.",
        "Provide comprehensive warranties and guarantees.",
    ]
    for r in recommendations:
        doc.add_paragraph(r)

    # Save document
    doc.save("CulturalAnalysisReport.docx")


create_cultural_analysis_report()
