from docx import Document


def create_research_objective_document():
    doc = Document()

    # Add title
    doc.add_heading("Research Objective Document", level=1)

    # Purpose of the Study
    doc.add_heading("Purpose of the Study", level=2)
    doc.add_paragraph(
        "To understand consumer preferences and cultural nuances in the target market."
    )

    # Key Questions to Address
    doc.add_heading("Key Questions to Address", level=2)
    questions = [
        "What specific information do we need to achieve our market traction goals?",
        "Which segments of the population should be prioritized for this study?",
    ]
    for q in questions:
        doc.add_paragraph(q)

    # Target Audience Segmentation
    doc.add_heading("Target Audience Segmentation", level=2)
    doc.add_paragraph(
        "Focus on urban professionals aged 25-40 with medium to high disposable income."
    )

    # Data Sources Identified
    doc.add_heading("Data Sources Identified", level=2)
    doc.add_paragraph("Industry reports, academic studies, and competitor analyses.")

    # Save document
    doc.save("ResearchObjectiveDocument.docx")


create_research_objective_document()
