import pandas as pd
from docx import Document


# Creating Excel Template for Stakeholder Matrix
def create_stakeholder_matrix():
    matrix_data = {
        "Stakeholder Name": [],
        "Category": [],
        "Influence (High/Medium/Low)": [],
        "Interest (High/Medium/Low)": [],
        "Potential Conflicts": [],
        "Alignment Opportunities": [],
    }

    df = pd.DataFrame(matrix_data)
    output_path = "/mnt/data/Stakeholder_Matrix_Template.xlsx"
    df.to_excel(output_path, index=False)
    return output_path


# Creating Word Template for Stakeholder List
def create_stakeholder_list():
    doc = Document()
    doc.add_heading("Stakeholder List", level=1)

    doc.add_paragraph("Purpose:")
    doc.add_paragraph(
        "To document all identified stakeholders, their categories, and relevant details during the identification process."
    )

    doc.add_heading("Stakeholder Categories", level=2)
    doc.add_paragraph("- Internal Stakeholders: [Example: Executives, Managers]")
    doc.add_paragraph("- External Stakeholders: [Example: Clients, Regulators]")

    doc.add_heading("Stakeholder Details", level=2)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Stakeholder Name"
    hdr_cells[1].text = "Category"
    hdr_cells[2].text = "Notes"

    output_path = "/mnt/data/Stakeholder_List_Template.docx"
    doc.save(output_path)
    return output_path


# Creating Word Template for Feedback Summary
def create_feedback_summary():
    doc = Document()
    doc.add_heading("Feedback Summary Report", level=1)

    doc.add_paragraph("Purpose:")
    doc.add_paragraph(
        "To consolidate and address feedback received from stakeholders during the stakeholder identification and mapping process."
    )

    doc.add_heading("Feedback Overview", level=2)
    doc.add_paragraph("- Stakeholder Name: [Example: John Doe]")
    doc.add_paragraph(
        "- Feedback Provided: [Example: Request for more detailed role definition]"
    )

    doc.add_heading("Actions Taken", level=2)
    doc.add_paragraph(
        "- Changes Implemented: [Example: Updated the role chart to reflect responsibilities]"
    )

    doc.add_heading("Remaining Concerns", level=2)
    doc.add_paragraph("[Example: Specific communication channels still unclear]")

    output_path = "/mnt/data/Feedback_Summary_Template.docx"
    doc.save(output_path)
    return output_path


# Creating Excel Template for Stakeholder Personas
def create_stakeholder_personas():
    persona_data = {
        "Stakeholder Name": [],
        "Role": [],
        "Objectives": [],
        "Concerns": [],
        "Preferred Communication Style": [],
        "Influence (High/Medium/Low)": [],
    }

    df = pd.DataFrame(persona_data)
    output_path = "/mnt/data/Stakeholder_Personas_Template.xlsx"
    df.to_excel(output_path, index=False)
    return output_path


# Creating Word Template for Conflict and Synergy Summary
def create_conflict_synergy_summary():
    doc = Document()
    doc.add_heading("Conflict and Synergy Summary", level=1)

    doc.add_paragraph("Purpose:")
    doc.add_paragraph(
        "To document identified conflicts and synergies among stakeholders during the mapping process."
    )

    doc.add_heading("Potential Conflicts", level=2)
    doc.add_paragraph(
        "- Stakeholder 1 vs. Stakeholder 2: [Example conflict description]"
    )

    doc.add_heading("Alignment Opportunities", level=2)
    doc.add_paragraph(
        "- Stakeholder 3 and Stakeholder 4: [Example synergy description]"
    )

    doc.add_heading("Proposed Mitigation Strategies", level=2)
    doc.add_paragraph("- Strategy for Conflict 1: [Example mitigation strategy]")

    output_path = "/mnt/data/Conflict_Synergy_Summary_Template.docx"
    doc.save(output_path)
    return output_path


# Creating Word Template for Finalized Stakeholder Map
def create_finalized_stakeholder_map():
    doc = Document()
    doc.add_heading("Finalized Stakeholder Map", level=1)

    doc.add_paragraph("Purpose:")
    doc.add_paragraph(
        "To present the validated and finalized stakeholder mapping document for alignment purposes."
    )

    doc.add_heading("Stakeholder Categories", level=2)
    doc.add_paragraph("- Internal Stakeholders: [Example: Executives, Managers]")
    doc.add_paragraph("- External Stakeholders: [Example: Clients, Regulators]")

    doc.add_heading("Validated Stakeholder Matrix", level=2)
    doc.add_paragraph("[Insert final stakeholder matrix here]")

    doc.add_heading("Sign-Offs", level=2)
    doc.add_paragraph("- Stakeholder Representative: [Name and Signature]")
    doc.add_paragraph("- Date: [Insert Date]")

    output_path = "/mnt/data/Finalized_Stakeholder_Map_Template.docx"
    doc.save(output_path)
    return output_path


# Generate templates
stakeholder_matrix_path = create_stakeholder_matrix()
stakeholder_list_path = create_stakeholder_list()
feedback_summary_path = create_feedback_summary()
stakeholder_personas_path = create_stakeholder_personas()
conflict_synergy_summary_path = create_conflict_synergy_summary()
finalized_stakeholder_map_path = create_finalized_stakeholder_map()

(
    stakeholder_matrix_path,
    stakeholder_list_path,
    feedback_summary_path,
    stakeholder_personas_path,
    conflict_synergy_summary_path,
    finalized_stakeholder_map_path,
)
