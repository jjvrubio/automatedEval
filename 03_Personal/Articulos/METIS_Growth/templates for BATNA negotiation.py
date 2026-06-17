from docx import Document
import pandas as pd


# Create a structured document template for the Sale Pattern Framework
def create_sale_pattern_framework():
    doc = Document()
    doc.add_heading("Sale Pattern Framework", 0)

    doc.add_heading("1. Optimized Sale Structure", level=1)
    doc.add_paragraph(
        "This section defines the structured and repeatable sales configuration "
        "that has been optimized based on prior deal analysis."
    )

    # Table of optimized sale elements
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Sale Element"
    hdr_cells[1].text = "Optimized Configuration"

    sale_elements = [
        ("Base Price", "$500 per unit"),
        ("Payment Terms", "30% upfront, 18-month installment"),
        ("Trial Period", "90-day trial with performance guarantee"),
        ("Warranty", "5 years standard"),
        ("Exclusivity", "1-year regional exclusivity for top-tier partners"),
    ]

    for element, value in sale_elements:
        row_cells = table.add_row().cells
        row_cells[0].text = element
        row_cells[1].text = value

    doc.add_heading("2. Standardized Contract Templates", level=1)
    doc.add_paragraph(
        "Predefined contract structures to ensure consistency across sales engagements."
    )

    doc.add_heading("3. Deal Success Playbook", level=1)
    doc.add_paragraph(
        "This section documents lessons learned from prior deals, including common objections and "
        "successful negotiation strategies."
    )

    doc.add_heading("4. Continuous Optimization Mechanism", level=1)
    doc.add_paragraph(
        "A structured feedback loop ensuring ongoing refinement of the Sale Pattern Framework."
    )

    doc_path = "/mnt/data/Sale_Pattern_Framework.docx"
    doc.save(doc_path)
    return doc_path


# Create an Excel sheet template for deal tracking
def create_deal_tracking_sheet():
    data = {
        "Client Name": [],
        "Industry": [],
        "Initial Objection": [],
        "Adjusted Offer": [],
        "Outcome": [],
    }
    df = pd.DataFrame(data)

    excel_path = "/mnt/data/Deal_Tracking_Sheet.xlsx"
    df.to_excel(excel_path, index=False)
    return excel_path


# Generate the documents
sale_pattern_path = create_sale_pattern_framework()
deal_tracking_path = create_deal_tracking_sheet()

# Provide download links
sale_pattern_path, deal_tracking_path
