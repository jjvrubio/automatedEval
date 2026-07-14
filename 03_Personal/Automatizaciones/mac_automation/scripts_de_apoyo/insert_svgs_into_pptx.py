#!/usr/bin/env python3
"""Inserta una secuencia de SVG como nuevas diapositivas en una presentación PPTX."""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.image import ImagePart


SLIDE_NUMBER_RE = re.compile(r"slide_(\d+)", re.IGNORECASE)
SVG_CONTENT_TYPE = "image/svg+xml"


class SvgImagePart(ImagePart):
    """ImagePart mínimo para SVG, sin depender de Pillow."""

    def __init__(self, partname, package, blob: bytes, filename: str):
        super().__init__(partname, SVG_CONTENT_TYPE, package, blob, filename)
        self._sha1 = hashlib.sha1(blob).hexdigest()

    @property
    def sha1(self) -> str:
        return self._sha1

    def scale(self, scaled_cx, scaled_cy):
        if scaled_cx is None or scaled_cy is None:
            raise ValueError("La inserción SVG requiere width y height explícitos.")
        return scaled_cx, scaled_cy


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inserta SVG como diapositivas nuevas al final de un archivo PPTX."
    )
    parser.add_argument("svg_source", help="Carpeta con SVG o archivo SVG individual.")
    parser.add_argument("pptx_path", help="Presentación PowerPoint de destino.")
    parser.add_argument(
        "--output",
        help="Ruta del PPTX de salida. Si no se indica, se crea una copia junto al original.",
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        help="Busca SVG de forma recursiva dentro de la carpeta indicada.",
    )
    parser.add_argument(
        "--pattern",
        default="*.svg",
        help="Patrón glob para filtrar SVG dentro de la carpeta. Por defecto: *.svg",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Sobrescribe el PPTX original en lugar de crear una copia.",
    )
    return parser.parse_args()


def slide_sort_key(path: Path) -> tuple[int, str]:
    match = SLIDE_NUMBER_RE.search(path.name)
    if match:
        return int(match.group(1)), path.name.lower()
    return sys.maxsize, path.name.lower()


def collect_svgs(source: Path, recursive: bool, pattern: str) -> list[Path]:
    if source.is_file():
        if source.suffix.lower() != ".svg":
            raise ValueError(f"El archivo no es SVG: {source}")
        return [source]

    if not source.is_dir():
        raise FileNotFoundError(f"No existe la ruta de SVG: {source}")

    glob_pattern = f"**/{pattern}" if recursive else pattern
    svgs = sorted(
        [path for path in source.glob(glob_pattern) if path.is_file() and path.suffix.lower() == ".svg"],
        key=slide_sort_key,
    )
    return svgs


def choose_blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
    return prs.slide_layouts[6]


def default_output_path(pptx_path: Path) -> Path:
    return pptx_path.with_name(f"{pptx_path.stem}-with-svg-slides{pptx_path.suffix}")


def get_or_add_svg_image_part(slide, svg_path: Path) -> tuple[SvgImagePart, str]:
    package = slide.part._package
    blob = svg_path.read_bytes()
    sha1 = hashlib.sha1(blob).hexdigest()

    for image_part in package._image_parts:
        if getattr(image_part, "sha1", None) == sha1:
            r_id = slide.part.relate_to(image_part, RT.IMAGE)
            return image_part, r_id

    image_part = SvgImagePart(package.next_image_partname("svg"), package, blob, svg_path.name)
    r_id = slide.part.relate_to(image_part, RT.IMAGE)
    return image_part, r_id


def add_svg_picture(slide, svg_path: Path, width: int, height: int) -> None:
    image_part, r_id = get_or_add_svg_image_part(slide, svg_path)
    slide.shapes._add_pic_from_image_part(image_part, r_id, 0, 0, width, height)
    slide.shapes._recalculate_extents()


def insert_svgs(svg_paths: list[Path], pptx_path: Path, output_path: Path) -> tuple[int, int]:
    prs = Presentation(pptx_path)
    initial_slide_count = len(prs.slides)
    blank_layout = choose_blank_layout(prs)
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    for svg_path in svg_paths:
        slide = prs.slides.add_slide(blank_layout)
        if svg_path.suffix.lower() == ".svg":
            add_svg_picture(slide, svg_path, slide_width, slide_height)
        else:
            slide.shapes.add_picture(str(svg_path), 0, 0, width=slide_width, height=slide_height)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return initial_slide_count, len(prs.slides)


def main() -> int:
    args = parse_args()

    svg_source = Path(args.svg_source).expanduser()
    pptx_path = Path(args.pptx_path).expanduser()
    if not pptx_path.exists():
        print(f"No existe el PPTX de destino: {pptx_path}", file=sys.stderr)
        return 1

    try:
        svg_paths = collect_svgs(svg_source, args.recursive, args.pattern)
    except (FileNotFoundError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 1

    if not svg_paths:
        print("No se encontraron SVG para insertar.", file=sys.stderr)
        return 1

    if args.overwrite:
        output_path = pptx_path
    elif args.output:
        output_path = Path(args.output).expanduser()
    else:
        output_path = default_output_path(pptx_path)

    try:
        initial_slide_count, final_slide_count = insert_svgs(svg_paths, pptx_path, output_path)
    except Exception as exc:
        print(f"No se pudo generar la presentación: {exc}", file=sys.stderr)
        return 1

    print(f"SVG insertados: {len(svg_paths)}")
    print(f"Diapositivas antes: {initial_slide_count}")
    print(f"Diapositivas después: {final_slide_count}")
    print(f"Archivo generado: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())