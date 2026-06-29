#!/usr/bin/env python3
"""Build a PPTX deck from SVG files, one SVG per slide."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.image import ImagePart
from pptx.util import Inches


SVG_CONTENT_TYPE = "image/svg+xml"
ORDINAL_RE = re.compile(r"^(\d+)")
LENGTH_RE = re.compile(r"^\s*([+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*([A-Za-z%]*)\s*$")


@dataclass(frozen=True)
class SvgInfo:
    path: str
    width: float | None
    height: float | None
    view_box: str | None
    paths: int
    groups: int
    texts: int
    images: int
    gradients: int
    patterns: int
    filters: int
    masks: int
    clip_paths: int
    foreign_objects: int


class SvgImagePart(ImagePart):
    """Minimal SVG ImagePart for python-pptx, avoiding Pillow decoding."""

    def __init__(self, partname, package, blob: bytes, filename: str):
        super().__init__(partname, SVG_CONTENT_TYPE, package, blob, filename)
        self._sha1 = hashlib.sha1(blob).hexdigest()

    @property
    def sha1(self) -> str:
        return self._sha1

    def scale(self, scaled_cx, scaled_cy):
        if scaled_cx is None or scaled_cy is None:
            raise ValueError("SVG insertion requires explicit width and height.")
        return scaled_cx, scaled_cy


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a PPTX deck from SVG files, using one full-slide SVG per slide."
    )
    parser.add_argument(
        "inputs",
        nargs="+",
        help="SVG files or directories containing SVG files.",
    )
    parser.add_argument(
        "--output",
        "-o",
        type=Path,
        default=Path("output/svg-deck.pptx"),
        help="PPTX output path. Default: output/svg-deck.pptx",
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        help="Search SVG files recursively inside input directories.",
    )
    parser.add_argument(
        "--pattern",
        default="*.svg",
        help="Glob pattern used for directory inputs. Default: *.svg",
    )
    parser.add_argument(
        "--template",
        type=Path,
        help="Optional PPTX template. New SVG slides are appended to it.",
    )
    parser.add_argument(
        "--slide-width",
        type=float,
        default=13.333333,
        help="Slide width in inches when no template is used. Default: 13.333333.",
    )
    parser.add_argument(
        "--slide-height",
        type=float,
        default=7.5,
        help="Slide height in inches when no template is used. Default: 7.5.",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        help="Optional JSON manifest with source order and SVG diagnostics.",
    )
    return parser.parse_args()


def ordinal_sort_key(path: Path) -> tuple[int, str]:
    match = ORDINAL_RE.match(path.name)
    if match:
        return int(match.group(1)), path.name.lower()
    return sys.maxsize, path.name.lower()


def collect_svgs(inputs: Iterable[str], recursive: bool, pattern: str) -> list[Path]:
    paths: list[Path] = []

    for raw_input in inputs:
        path = Path(raw_input).expanduser()
        if path.is_file():
            if path.suffix.lower() != ".svg":
                raise ValueError(f"Not an SVG file: {path}")
            paths.append(path)
            continue

        if path.is_dir():
            glob_pattern = f"**/{pattern}" if recursive else pattern
            paths.extend(
                child
                for child in path.glob(glob_pattern)
                if child.is_file() and child.suffix.lower() == ".svg"
            )
            continue

        raise FileNotFoundError(f"Input path does not exist: {path}")

    unique_paths = list(dict.fromkeys(paths))
    return sorted(unique_paths, key=ordinal_sort_key)


def strip_ns(tag: str) -> str:
    if tag.startswith("{"):
        return tag.rsplit("}", 1)[1]
    return tag


def parse_numeric_length(raw_value: str | None) -> float | None:
    if raw_value is None:
        return None
    match = LENGTH_RE.match(raw_value)
    if not match:
        return None
    unit = match.group(2)
    if unit not in {"", "px"}:
        return None
    return float(match.group(1))


def inspect_svg(path: Path) -> SvgInfo:
    root = ET.parse(path).getroot()
    counts: dict[str, int] = {
        "path": 0,
        "g": 0,
        "text": 0,
        "image": 0,
        "linearGradient": 0,
        "radialGradient": 0,
        "pattern": 0,
        "filter": 0,
        "mask": 0,
        "clipPath": 0,
        "foreignObject": 0,
    }

    for element in root.iter():
        tag = strip_ns(element.tag)
        if tag in counts:
            counts[tag] += 1

    return SvgInfo(
        path=str(path),
        width=parse_numeric_length(root.get("width")),
        height=parse_numeric_length(root.get("height")),
        view_box=root.get("viewBox"),
        paths=counts["path"],
        groups=counts["g"],
        texts=counts["text"],
        images=counts["image"],
        gradients=counts["linearGradient"] + counts["radialGradient"],
        patterns=counts["pattern"],
        filters=counts["filter"],
        masks=counts["mask"],
        clip_paths=counts["clipPath"],
        foreign_objects=counts["foreignObject"],
    )


def choose_blank_layout(prs: Presentation):
    for layout in prs.slide_layouts:
        if len(layout.placeholders) == 0:
            return layout
    return prs.slide_layouts[6]


def build_presentation(template: Path | None, slide_width: float, slide_height: float) -> Presentation:
    prs = Presentation(template) if template else Presentation()
    if template is None:
        prs.slide_width = Inches(slide_width)
        prs.slide_height = Inches(slide_height)
    return prs


def get_or_add_svg_image_part(slide, svg_path: Path) -> tuple[SvgImagePart, str]:
    package = slide.part._package
    blob = svg_path.read_bytes()
    sha1 = hashlib.sha1(blob).hexdigest()

    for image_part in package._image_parts:
        if getattr(image_part, "sha1", None) == sha1:
            r_id = slide.part.relate_to(image_part, RT.IMAGE)
            return image_part, r_id

    image_part = SvgImagePart(package.next_image_partname("svg"), package, blob, svg_path.name)
    r_id = slide.part.relate_to(image_part, RT.IMAGE)
    return image_part, r_id


def add_full_slide_svg(slide, svg_path: Path, width: int, height: int) -> None:
    image_part, r_id = get_or_add_svg_image_part(slide, svg_path)
    picture = slide.shapes._add_pic_from_image_part(image_part, r_id, 0, 0, width, height)
    picture.name = svg_path.stem
    slide.shapes._recalculate_extents()


def write_manifest(path: Path, infos: list[SvgInfo], output_path: Path) -> None:
    payload = {
        "output": str(output_path),
        "slide_count": len(infos),
        "slides": [asdict(info) for info in infos],
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def create_deck(
    svg_paths: list[Path],
    output_path: Path,
    template: Path | None,
    slide_width: float,
    slide_height: float,
) -> list[SvgInfo]:
    prs = build_presentation(template, slide_width, slide_height)
    blank_layout = choose_blank_layout(prs)
    infos: list[SvgInfo] = []

    for svg_path in svg_paths:
        infos.append(inspect_svg(svg_path))
        slide = prs.slides.add_slide(blank_layout)
        add_full_slide_svg(slide, svg_path, prs.slide_width, prs.slide_height)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return infos


def main() -> int:
    args = parse_args()

    try:
        svg_paths = collect_svgs(args.inputs, args.recursive, args.pattern)
    except (FileNotFoundError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 1

    if not svg_paths:
        print("No SVG files found.", file=sys.stderr)
        return 1

    if args.template and not args.template.expanduser().exists():
        print(f"Template does not exist: {args.template}", file=sys.stderr)
        return 1

    output_path = args.output.expanduser()
    manifest_path = args.manifest.expanduser() if args.manifest else output_path.with_suffix(".manifest.json")

    try:
        infos = create_deck(
            svg_paths=svg_paths,
            output_path=output_path,
            template=args.template.expanduser() if args.template else None,
            slide_width=args.slide_width,
            slide_height=args.slide_height,
        )
        write_manifest(manifest_path, infos, output_path)
    except Exception as exc:
        print(f"Could not create PPTX: {exc}", file=sys.stderr)
        return 1

    print(f"SVG files: {len(svg_paths)}")
    print(f"PPTX: {output_path}")
    print(f"Manifest: {manifest_path}")
    for index, svg_path in enumerate(svg_paths, start=1):
        print(f"{index:02d}. {svg_path.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
