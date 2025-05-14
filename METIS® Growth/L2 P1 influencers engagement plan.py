from docx import Document

def create_influencer_engagement_plan():
    doc = Document()
    
    # Add title
    doc.add_heading('Influencer Engagement Plan', level=1)
    
    # List of Potential Partners
    doc.add_heading('List of Potential Partners', level=2)
    influencers = [
        {'Name': 'Jane Doe', 'Followers': '500k', 'Topics': 'Health & Wellness'},
        {'Name': 'John Smith', 'Followers': '300k', 'Topics': 'Tech & Gadgets'}
    ]
    for influencer in influencers:
        doc.add_paragraph(f"Name: {influencer['Name']}, Followers: {influencer['Followers']}, Topics: {influencer['Topics']}")
    
    # Proposed Collaboration Strategies
    doc.add_heading('Proposed Collaboration Strategies', level=2)
    strategies = [
        'Co-create wellness-themed posts.',
        'Offer exclusive discount codes.'
    ]
    for s in strategies:
        doc.add_paragraph(s)
    
    # Performance Metrics
    doc.add_heading('Performance Metrics', level=2)
    doc.add_paragraph('Track engagement rate, reach, and conversion rates.')
    
    # Save document
    doc.save('InfluencerEngagementPlan.docx')

create_influencer_engagement_plan()