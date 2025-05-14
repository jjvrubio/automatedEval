import os
import docx
import PyPDF2
import anthropic
import json
import logging

# Configure logging
logging.basicConfig(
    filename="grading_debug.log",
    filemode="a",
    format="%(asctime)s - %(levelname)s - %(message)s",
    level=logging.DEBUG,
)
logging.info("Logging initialized. Starting script...")

class FileReader:
    """Handles reading and cleaning of file content."""

    @staticmethod
    def clean_text(text: str) -> str:
        """Clean and validate extracted text."""
        if not text.strip():
            raise ValueError("Extracted text is empty.")
        return ''.join(char for char in text if ord(char) >= 32 or char in '\n\r\t').strip()

    @staticmethod
    def read_docx(file_path: str) -> str:
        """Read content from a DOCX file."""
        try:
            doc = docx.Document(file_path)
            return FileReader.clean_text('\n'.join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()))
        except Exception as e:
            raise Exception(f"Error reading DOCX file {file_path}: {str(e)}")

    @staticmethod
    def read_pdf(file_path: str) -> str:
        """Read content from a PDF file."""
        try:
            with open(file_path, 'rb') as pdf_file:
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = [page.extract_text() for page in pdf_reader.pages]
                return FileReader.clean_text('\n'.join(text))
        except Exception as e:
            raise Exception(f"Error reading PDF file {file_path}: {str(e)}")

    @staticmethod
    def read_file(file_path: str) -> str:
        """Read and clean content from various file types."""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".docx":
            return FileReader.read_docx(file_path)
        elif ext == ".pdf":
            return FileReader.read_pdf(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")


class AnthropicClient:
    """Handles communication with the Anthropic API."""

    def __init__(self, api_key: str):
        try:
            self.client = anthropic.Client(api_key=api_key)
        except Exception as e:
            raise Exception(f"Error initializing Anthropic client: {str(e)}")

    def get_completion(self, prompt: str, model: str, max_tokens: int = 4000, temperature: float = 0.0) -> str:
        """Fetch completion from Anthropic Messages API."""
        response = self.client.messages.create(
            model=model,
            max_tokens=max_tokens,
            temperature=temperature,
            messages=[{"role": "user", "content": prompt}],
        )
        if response.content:
            return response.content[0].text
        raise ValueError("API returned an empty response.")


class DocumentGrader:
    """Coordinates grading using file reading and Anthropic API."""

    def __init__(self, api_client: AnthropicClient):
        self.api_client = api_client

    def prepare_grading_prompt(self, exercise: str, rubric: str, solution: str) -> str:
        """Prepare the prompt for Anthropic API."""
        return f"""\n\nHuman: Please grade the following student solution based on the exercise description and grading rubric provided.

Exercise Description:
{exercise}

Grading Rubric:
{rubric}

Student's Solution:
{solution}

Please provide your response in valid JSON format:
{{
    "analysis": "Your analysis here",
    "criterion_scores": {{
        "criterion1": "score",
        "criterion2": "score"
    }},
    "total_score": "final_score",
    "feedback": "Your feedback here"
}}
\n\nAssistant:"""

    def grade_solution(self, exercise_path: str, rubric_path: str, solution_path: str) -> dict:
        """Grade the solution using Anthropic API."""
        exercise = FileReader.read_file(exercise_path)
        rubric = FileReader.read_file(rubric_path)
        solution = FileReader.read_file(solution_path)

        logging.info(f"Prepared prompt for solution: {solution_path}")
        prompt = self.prepare_grading_prompt(exercise, rubric, solution)

        response_text = self.api_client.get_completion(prompt, model="claude-3-haiku-20240307")
        return json.loads(response_text)

    def save_results(self, results: dict, output_path: str, exclude_keys=None):
        """Save grading results to a Word document and JSON file."""
        exclude_keys = exclude_keys or []
        # Save to JSON
        json_output = f"{os.path.splitext(output_path)[0]}.json"
        trimmed_results = {k: v for k, v in results.items() if k not in exclude_keys}
        with open(json_output, "w") as json_file:
            json.dump(trimmed_results, json_file, indent=4)
        logging.info(f"Results saved to JSON: {json_output}")

        # Save to DOCX
        doc = docx.Document()
        doc.add_heading("Grading Results", 0)

        for section, key in [("Analysis", "analysis"), ("Total Score", "total_score"), ("Feedback", "feedback")]:
            doc.add_heading(section, level=1)
            doc.add_paragraph(str(results.get(key, f"No {section.lower()} provided")))

        doc.add_heading("Criterion Scores", level=1)
        for criterion, score in results.get("criterion_scores", {}).items():
            doc.add_paragraph(f"{criterion}: {score}")

        doc.save(output_path)
        logging.info(f"Results saved to DOCX: {output_path}")


def main():
    try:
        api_key = os.getenv("MI_CLAVE_API_ANTHROPIC")
        if not api_key:
            raise ValueError("Anthropic API key not found in environment variables.")

        grader = DocumentGrader(AnthropicClient(api_key))

        exercise_path = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Exercise instructions.docx"
        rubric_path = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/interview rubric.docx"
        solutions_dir = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Class exam/"
        output_dir = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Class exam graded/"

        os.makedirs(output_dir, exist_ok=True)

        for solution_file in os.listdir(solutions_dir):
            solution_path = os.path.join(solutions_dir, solution_file)

            if not solution_file.lower().endswith((".docx", ".pdf")):
                logging.warning(f"Skipping unsupported file: {solution_file}")
                continue

            try:
                logging.info(f"Processing file: {solution_file}")
                results = grader.grade_solution(exercise_path, rubric_path, solution_path)

                # Save outputs
                output_path = os.path.join(output_dir, f"{os.path.splitext(solution_file)[0]}_graded.docx")
                grader.save_results(results, output_path, exclude_keys=["analysis", "feedback"])
            except Exception as e:
                logging.error(f"Error grading {solution_file}: {str(e)}")

        logging.info("Grading completed successfully.")
    except Exception as e:
        logging.error(f"Error: {e}")


if __name__ == "__main__":
    main()
