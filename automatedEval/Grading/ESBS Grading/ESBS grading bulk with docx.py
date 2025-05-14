# Cambiar los paths de los archivos a leer y escribir

from AppKit import NSOpenPanel, NSApplication, NSApp
from Foundation import NSURL

import os
import json
import docx
import markdown
from openai import OpenAI
from typing import Dict, List


class DocumentGrader:
    def __init__(self, api_key: str):
        self.client = OpenAI(api_key=api_key)

    def read_docx(self, file_path: str) -> str:
        """Read content from a Word document."""
        try:
            doc = docx.Document(file_path)
            return '\n'.join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
        except Exception as e:
            raise Exception(f"Error reading DOCX document {file_path}: {e}")

    def read_markdown(self, file_path: str) -> str:
        """Read and parse content from a Markdown file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as md_file:
                return markdown.markdown(md_file.read())
        except Exception as e:
            raise Exception(f"Error reading Markdown document {file_path}: {e}")

    def prepare_grading_prompt(self, exercise: str, rubric: str, solution: str) -> str:
        """Prepare the grading prompt for OpenAI."""
        return f"""Please grade the following student solution based on the exercise description and grading rubric provided.

        Exercise Description:
        {exercise}

        Grading Rubric:
        {rubric}

        Student's Solution:
        {solution}

        Please provide:
        1. A detailed analysis of the solution
        2. Points awarded for each rubric criterion
        3. Total score
        4. Specific feedback and suggestions for improvement

        Format your response ONLY as JSON with the following structure:
        {{
            "analysis": "detailed analysis here",
            "criterion_scores": {{"criterion1": "score", "criterion2": "score"}},
            "total_score": "numerical_score",
            "feedback": "detailed feedback here"
        }}"""

    def grade_solution(self, exercise_path: str, rubric_path: str, solution_path: str) -> Dict:
        """Grade the solution using OpenAI API."""
        try:
            exercise = self.read_markdown(exercise_path)
            rubric = self.read_markdown(rubric_path)
            solution = self.read_docx(solution_path)
            prompt = self.prepare_grading_prompt(exercise, rubric, solution)    

            response = self.client.chat.completions.create(
                model="gpt-4-turbo",
                messages=[{"role": "system", "content": "You are an expert grader."},
                          {"role": "user", "content": prompt}],
                temperature=0.0,
                max_tokens=4000,
            )
            return json.loads(response.choices[0].message.content.strip())
        except Exception as e:
            raise Exception(f"Grading failed: {e}")

    def save_results(self, results: Dict, output_base_path: str):
        """Save results to DOCX and JSON files."""
        try:
            # Save to DOCX
            doc = docx.Document()
            doc.add_heading('Grading Results', 0).alignment = 1
            doc.add_heading('Analysis', level=1)
            doc.add_paragraph(results.get('analysis', 'No analysis provided'))
            doc.add_heading('Criterion Scores', level=1)
            for criterion, score in results.get('criterion_scores', {}).items():
                doc.add_paragraph(f"{criterion}: {score}")
            doc.add_heading('Total Score', level=1)
            doc.add_paragraph(str(results.get('total_score', 'No score available')))
            doc.add_heading('Feedback', level=1)
            doc.add_paragraph(results.get('feedback', 'No feedback provided'))
            doc.save(f"{output_base_path}.docx")

            # Save to JSON
            with open(f"{output_base_path}.json", 'w', encoding='utf-8') as json_file:
                json.dump(results, json_file, indent=4)
            print(f"Results saved to: {output_base_path}.docx and {output_base_path}.json")

        except Exception as e:
            raise Exception(f"Error saving results: {e}")

def select_file(allowed_file_types=None, message="Select a file"):
    panel = NSOpenPanel.openPanel()
    panel.setTitle_(message)
    panel.setCanChooseFiles_(True)
    panel.setCanChooseDirectories_(False)
    panel.setAllowsMultipleSelection_(False)
    if allowed_file_types:
        panel.setAllowedFileTypes_(allowed_file_types)
    if panel.runModal():
        return str(panel.URLs()[0].path())
    return None

def select_folder(message="Select a folder"):
    panel = NSOpenPanel.openPanel()
    panel.setTitle_(message)
    panel.setCanChooseFiles_(False)
    panel.setCanChooseDirectories_(True)
    panel.setAllowsMultipleSelection_(False)
    if panel.runModal():
        return str(panel.URLs()[0].path())
    return None

def prepare_grading_prompt(self, exercise: str, rubric: str, solution: str) -> str:
    """Prepare a strict grading prompt based on the rubric and instructions."""
    return f"""You are an academic evaluator. You must grade the student's solution according to the provided rubric and the detailed exercise instructions.

🚫 DO NOT invent, soften, or ignore the rubric structure.
✅ USE ONLY the defined heuristics and point ranges per rubric category, which are structured in XML format. 
📋 The evaluation must reflect each category separately and assign a score within the allowed values: 0, 1–3, 4–5, 6–7, 8–9, or 10. Never assign midpoints or undefined scores.

💡 The structure of the assignment is explicitly defined within <toc> tags in the exercise instructions. You must verify that the student addresses all required sections.

🛠️ Your evaluation must follow this strict structure:
1. An individual analysis per rubric category (Argument Quality, Clarity, Coherence, Grammar, Use of Evidence, Syntax).
2. Points awarded per rubric category.
3. A clear explanation of how each point was determined based on the rubric.
4. Total numerical score (sum or weighted average if needed).
5. Actionable feedback for improvement.
6. Identify missing sections if any required parts (e.g., performance tables, forecast, APA references) are absent.

📂 EXERCISE INSTRUCTIONS:
{exercise}

📏 GRADING RUBRIC (XML categories with heuristics):
{rubric}

📝 STUDENT SOLUTION:
{solution}

🎯 FORMAT your full response as JSON using this schema:
{{
    "analysis": {{
        "Argument Quality": "...",
        "Clarity": "...",
        "Coherence": "...",
        "Grammar": "...",
        "Use of Evidence": "...",
        "Syntax": "..."
    }},
    "criterion_scores": {{
        "Argument Quality": score,
        "Clarity": score,
        "Coherence": score,
        "Grammar": score,
        "Use of Evidence": score,
        "Syntax": score
    }},
    "total_score": score,
    "feedback": "Detailed and actionable feedback, mentioning any missing structural elements or rubric violations."
}}"""



def main():
    try:
        # Get API key
        api_key = os.getenv('MI_CLAVE_API_OPENAI')
        if not api_key:
            raise ValueError("OpenAI API key not found in environment variables.")

        grader = DocumentGrader(api_key)

        # Paths form static folders version. I comment them just in case a backtrack.
        # exercise_path = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Exercise instructions.md"
        # rubric_path = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Grading rubric.md"
        # solution_dir = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Class exam"
        # output_dir = "/Users/juanjo/Library/CloudStorage/OneDrive-KALEIDAGEOGRAFIAS&MERCADOSSL/ESBS/Asignaturas/Consumer Market Research/Exercises/5th day interview/Class exam graded"
        
        # Native macOS dialogs
        print("Select the exercise Markdown file:")
        exercise_path = select_file(allowed_file_types=["md"], message="Select Exercise File")
        if not exercise_path:
            raise Exception("No exercise file selected.")

        print("Select the grading rubric Markdown file:")
        rubric_path = select_file(allowed_file_types=["md"], message="Select Rubric File")
        if not rubric_path:
            raise Exception("No rubric file selected.")

        print("Select the folder containing student solutions (DOCX):")
        solution_dir = select_folder("Select Solutions Folder")
        if not solution_dir:
            raise Exception("No solution directory selected.")

        print("Select the output folder for graded results:")
        output_dir = select_folder("Select Output Folder")
        if not output_dir:
            raise Exception("No output directory selected.")

        os.makedirs(output_dir, exist_ok=True)

        # Process each solution
        for solution_file in os.listdir(solution_dir):
            if solution_file.endswith('.docx'):
                solution_path = os.path.join(solution_dir, solution_file)
                output_base_path = os.path.join(output_dir, f"graded_{os.path.splitext(solution_file)[0]}")
                print(f"Grading solution: {solution_file}")
                results = grader.grade_solution(exercise_path, rubric_path, solution_path)
                grader.save_results(results, output_base_path)

        print("Grading completed for all solutions.")

    except Exception as e:
        print(f"Error: {e}")
        exit(1)


if __name__ == "__main__":
    main()
