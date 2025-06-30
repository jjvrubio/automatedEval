import openai
import json
import os
import re
import docx
import pdfplumber
import zipfile
import xml.etree.ElementTree as ET
from AppKit import NSApplication
from Cocoa import NSOpenPanel

def select_file(allowed_types, title):
    NSApplication.sharedApplication()
    panel = NSOpenPanel.openPanel()
    panel.setAllowedFileTypes_(allowed_types)
    panel.setCanChooseFiles_(True)
    panel.setCanChooseDirectories_(False)
    panel.setTitle_(title)
    if panel.runModal() == 1:
        return panel.URLs()[0].path()
    return None

def extract_text_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        text = ""
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    text_parts = []
    text_parts.extend(para.text for para in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    text_parts.append(cell_text)
    for section in doc.sections:
        header = section.header
        footer = section.footer
        if header:
            text_parts.extend(para.text for para in header.paragraphs if para.text.strip())
        if footer:
            text_parts.extend(para.text for para in footer.paragraphs if para.text.strip())
    return "\n".join(text_parts)

def find_first_line_with_anex(lines):
    for i, line in enumerate(lines):
        norm_line = line.lower()
        norm_line = norm_line.replace('á', 'a').replace('é', 'e').replace('í', 'i').replace('ó', 'o').replace('ú', 'u')
        if 'anex' in norm_line:
            return i
    return None

def extract_references_section(text):
    lines = text.splitlines()
    ref_indices = []
    for i, line in enumerate(lines):
        if re.search(r"REFERENCIAS|REFERENCIAS BIBLIOGRÁFICAS|BIBLIOGRAF[IÍ]A", line, re.IGNORECASE):
            ref_indices.append(i)
    if not ref_indices:
        return []
    ref_start = ref_indices[-1]
    ref_end = find_first_line_with_anex(lines[ref_start+1:])
    if ref_end is not None:
        ref_end = ref_start + 1 + ref_end
    return lines[ref_start+1:ref_end] if ref_end is not None else lines[ref_start+1:]

def get_paragraphs_with_terms_from_docx_xml(docx_path, terms):
    def normalize(s):
        return s.lower().replace('á','a').replace('é','e').replace('í','i').replace('ó','o').replace('ú','u')
    with zipfile.ZipFile(docx_path) as docx_zip:
        with docx_zip.open('word/document.xml') as xml_file:
            tree = ET.parse(xml_file)
            root = tree.getroot()
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            results = []
            idx = 0
            for para in root.findall('.//w:p', ns):
                texts = []
                for t in para.findall('.//w:t', ns):
                    texts.append(t.text or '')
                full_text = ''.join(texts)
                norm_text = normalize(full_text)
                if any(term in norm_text for term in terms):
                    results.append((idx, full_text))
                for instr in para.findall('.//w:instrText', ns):
                    if instr.text:
                        norm_instr = normalize(instr.text)
                        if any(term in norm_instr for term in terms):
                            results.append((idx, instr.text))
                idx += 1
            return results

def extract_references_section_docx_xml(docx_path):
    ref_terms = ['referencias', 'referencias bibliograficas', 'bibliografia']
    anex_terms = ['anex']
    ref_paras = get_paragraphs_with_terms_from_docx_xml(docx_path, ref_terms)
    anex_paras = get_paragraphs_with_terms_from_docx_xml(docx_path, anex_terms)
    if not ref_paras:
        return []
    ref_start_idx = ref_paras[-1][0]
    anex_after_ref = [idx for idx, txt in anex_paras if idx > ref_start_idx]
    if anex_after_ref:
        ref_end_idx = anex_after_ref[0]
    else:
        ref_end_idx = None
    with zipfile.ZipFile(docx_path) as docx_zip:
        with docx_zip.open('word/document.xml') as xml_file:
            tree = ET.parse(xml_file)
            root = tree.getroot()
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            all_paras = []
            for para in root.findall('.//w:p', ns):
                texts = []
                for t in para.findall('.//w:t', ns):
                    texts.append(t.text or '')
                full_text = ''.join(texts)
                all_paras.append(full_text)
            if ref_end_idx is not None:
                return all_paras[ref_start_idx+1:ref_end_idx]
            else:
                return all_paras[ref_start_idx+1:]

def join_reference_lines(ref_lines, max_length=500):
    joined = []
    buffer = ""
    ref_start_pattern = re.compile(r"^[A-ZÁÉÍÓÚÑ][a-záéíóúñA-ZÁÉÍÓÚÑ0-9 .,&-]+\(\d{4}")
    for line in ref_lines:
        if ref_start_pattern.match(line.strip()):
            if buffer:
                joined.append(buffer.strip())
            buffer = line.strip()
        else:
            buffer += " " + line.strip()
    if buffer:
        joined.append(buffer.strip())
    return joined

def main():
    api_key = os.environ.get("MI_CLAVE_API_OPENAI")
    if not api_key:
        print("ERROR: No se encontró la variable de entorno MI_CLAVE_API_OPENAI. Por favor, expórtala antes de ejecutar el script.")
        return
    tfm_file = select_file(["pdf", "docx"], "Selecciona el archivo del TFM (PDF o DOCX)")
    if not tfm_file:
        print("No se seleccionó archivo de TFM.")
        return
    ext = os.path.splitext(tfm_file)[-1].lower()
    if ext == ".pdf":
        cuerpo = extract_text_from_pdf(tfm_file)
        ref_lines = extract_references_section(cuerpo)
    elif ext == ".docx":
        ref_lines = extract_references_section_docx_xml(tfm_file)
        cuerpo = extract_text_from_docx(tfm_file)
    else:
        print("Formato no soportado.")
        return
    referencias = join_reference_lines(ref_lines)
    print(f"Referencias APA extraídas: {len(referencias)}")
    print(f"Ejemplo de referencia: {referencias[0] if referencias else 'Ninguna'}")
    prompt = f"""
Dado el siguiente cuerpo de texto académico y la lista de referencias bibliográficas en formato APA 7, realiza lo siguiente:
1. Extrae todas las citas en formato APA (por ejemplo, (Apellido, año), (Apellido1 & Apellido2, año), (Apellido et al., año), etc.) que aparecen en el texto.
2. Indica para cada cita si existe una referencia correspondiente en la lista de referencias (por autor y año).
3. Indica para cada referencia si está citada en el texto.
4. Si detectas errores de formato o inconsistencias, indícalo.
\nTexto del TFM:\n"""
    prompt += cuerpo[:12000]
    prompt += "\n\nLista de referencias APA:\n"
    prompt += '\n'.join(referencias)
    client = openai.OpenAI(api_key=api_key)
    response = client.chat.completions.create(
        model="gpt-4",
        messages=[
            {"role": "system", "content": "Eres un experto en normas APA y revisión académica."},
            {"role": "user", "content": prompt}
        ],
        max_tokens=1500,
        temperature=0.2
    )
    resultado = response.choices[0].message.content
    print("\n--- RESULTADO GPT-4 ---\n")
    print(resultado)
    # Guardar como Markdown
    out_dir = os.path.dirname(tfm_file)
    out_base = os.path.splitext(os.path.basename(tfm_file))[0]
    md_path = os.path.join(out_dir, f"{out_base}_informe_apa_gpt4.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# Informe de validación APA (GPT-4)\n\n")
        f.write(resultado)
    print(f"\nInforme guardado en: {md_path}\n")

if __name__ == "__main__":
    main()
